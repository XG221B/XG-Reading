#!/usr/bin/env python3
"""Build a Figure Evidence paper reader from a structured UTF-8 JSON file."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import tempfile
from contextlib import nullcontext
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont, ImageOps

SKILL_DIR = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = SKILL_DIR / "assets" / "reference.docx"
EXPECTED_TEMPLATE_SHA256 = (
    "A1EE0C4F7B46F20A8CA3EBFFE2224157E09FA0929A1F1DC02EC822CF6E2F8BB8"
)

CN_FONT = "华文中宋"
LATIN_FONT = "Times New Roman"
BLACK = RGBColor(0, 0, 0)

STYLE_TITLE = "FER Title"
STYLE_METADATA = "FER Metadata"
STYLE_SECTION = "FER Section"
STYLE_STEP = "FER Step"
STYLE_BODY = "FER Body"
STYLE_FIGURE = "FER Figure"
STYLE_FIGURE_TRANSLATION = "FER Figure Translation"
STYLE_SUPPLEMENT = "FER Supplement"
STYLE_PICTURE = "FER Picture"
STYLE_BULLET = "FER Bullet"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--content", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument(
        "--work-dir",
        type=Path,
        help="Optional durable directory for generated support-image sheets",
    )
    return parser.parse_args()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def clean_text(value: Any) -> str:
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    cjk = r"\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff"
    text = re.sub(rf"([{cjk}])\s+([A-Za-z0-9])", r"\1\2", text)
    text = re.sub(rf"([A-Za-z0-9])\s+([{cjk}])", r"\1\2", text)
    return text


def require_mapping(data: dict[str, Any], key: str) -> dict[str, Any]:
    value = data.get(key)
    if not isinstance(value, dict):
        raise ValueError(f"{key} must be an object")
    return value


def require_list(data: dict[str, Any], key: str) -> list[Any]:
    value = data.get(key)
    if not isinstance(value, list):
        raise ValueError(f"{key} must be an array")
    return value


def validate_content(data: dict[str, Any]) -> None:
    for key in ("document_title", "citation"):
        if not clean_text(data.get(key)):
            raise ValueError(f"{key} is required")

    positioning = require_mapping(data, "field_positioning")
    if not require_list(positioning, "paragraphs"):
        raise ValueError("field_positioning.paragraphs must not be empty")

    logic_chain = require_list(data, "logic_chain")
    if not logic_chain:
        raise ValueError("logic_chain must not be empty")
    for index, step_data in enumerate(logic_chain, start=1):
        if not isinstance(step_data, dict):
            raise ValueError(f"logic_chain[{index}] must be an object")
        for key in ("title", "text"):
            if not clean_text(step_data.get(key)):
                raise ValueError(f"logic_chain[{index}].{key} is required")

    blocks = data.get("evidence_blocks")
    if blocks is None:
        blocks = data.get("figures")
    if not isinstance(blocks, list) or not blocks:
        raise ValueError("evidence_blocks must be a non-empty array")

    for index, block in enumerate(blocks, start=1):
        if not isinstance(block, dict):
            raise ValueError(f"evidence_blocks[{index}] must be an object")
        for key in ("title", "image", "paper_claim", "logic_role"):
            if not clean_text(block.get(key)):
                raise ValueError(f"evidence_blocks[{index}].{key} is required")
        panels = block.get("panel_explanations")
        if not isinstance(panels, list) or not panels:
            raise ValueError(
                f"evidence_blocks[{index}].panel_explanations must not be empty"
            )
        for panel_index, panel in enumerate(panels, start=1):
            if isinstance(panel, dict):
                for key in ("title", "text"):
                    if not clean_text(panel.get(key)):
                        raise ValueError(
                            f"evidence_blocks[{index}].panel_explanations[{panel_index}].{key} is required"
                        )
            elif not clean_text(panel):
                raise ValueError(
                    f"evidence_blocks[{index}].panel_explanations[{panel_index}] must not be empty"
                )


def resolve_path(raw: str, base_dir: Path) -> Path:
    path = Path(raw)
    if not path.is_absolute():
        path = base_dir / path
    path = path.resolve()
    if not path.exists():
        raise FileNotFoundError(path)
    return path


def clear_document_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_style_font(style, size_pt: float, bold: bool = False) -> None:
    style.font.name = LATIN_FONT
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    style.font.color.rgb = BLACK
    rpr = style._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), LATIN_FONT)
    rfonts.set(qn("w:hAnsi"), LATIN_FONT)
    rfonts.set(qn("w:eastAsia"), CN_FONT)
    rfonts.set(qn("w:cs"), LATIN_FONT)


def get_or_add_style(document: Document, name: str):
    if name in document.styles:
        return document.styles[name]
    style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = document.styles["Normal"]
    return style


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    set_style_font(normal, 11)

    title = get_or_add_style(document, STYLE_TITLE)
    set_style_font(title, 16, True)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.line_spacing = 1.25
    title.paragraph_format.space_after = Pt(8)

    metadata = get_or_add_style(document, STYLE_METADATA)
    set_style_font(metadata, 11)
    metadata.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.line_spacing = 1.25
    metadata.paragraph_format.space_after = Pt(12)

    section = get_or_add_style(document, STYLE_SECTION)
    set_style_font(section, 13, True)
    section.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    section.paragraph_format.line_spacing = 1.25
    section.paragraph_format.space_before = Pt(10)
    section.paragraph_format.space_after = Pt(8)
    section.paragraph_format.keep_with_next = True
    section.paragraph_format.widow_control = True

    step = get_or_add_style(document, STYLE_STEP)
    set_style_font(step, 11, True)
    step.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    step.paragraph_format.line_spacing = 1.25
    step.paragraph_format.space_after = Pt(3)
    step.paragraph_format.keep_with_next = True
    step.paragraph_format.widow_control = True

    body = get_or_add_style(document, STYLE_BODY)
    set_style_font(body, 11)
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.first_line_indent = Cm(0.74)
    body.paragraph_format.line_spacing = 1.25
    body.paragraph_format.space_after = Pt(6)
    body.paragraph_format.widow_control = True

    figure = get_or_add_style(document, STYLE_FIGURE)
    set_style_font(figure, 13, True)
    figure.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure.paragraph_format.line_spacing = 1.25
    figure.paragraph_format.space_after = Pt(8)
    # Each main-figure evidence block begins on a fresh page.
    figure.paragraph_format.page_break_before = True
    figure.paragraph_format.keep_with_next = True
    figure.paragraph_format.widow_control = True

    figure_translation = get_or_add_style(document, STYLE_FIGURE_TRANSLATION)
    set_style_font(figure_translation, 11, True)
    figure_translation.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure_translation.paragraph_format.line_spacing = 1.25
    figure_translation.paragraph_format.space_before = Pt(6)
    figure_translation.paragraph_format.space_after = Pt(4)
    figure_translation.paragraph_format.keep_with_next = True
    figure_translation.paragraph_format.widow_control = True

    supplement = get_or_add_style(document, STYLE_SUPPLEMENT)
    set_style_font(supplement, 11, True)
    supplement.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    supplement.paragraph_format.line_spacing = 1.25
    supplement.paragraph_format.space_after = Pt(2)
    supplement.paragraph_format.keep_with_next = True

    picture = get_or_add_style(document, STYLE_PICTURE)
    set_style_font(picture, 11)
    picture.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.space_after = Pt(4)

    bullet = get_or_add_style(document, STYLE_BULLET)
    if "List Bullet" not in document.styles:
        raise ValueError("Retained template does not contain the List Bullet style")
    bullet.base_style = document.styles["List Bullet"]
    set_style_font(bullet, 11)
    bullet.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    bullet.paragraph_format.left_indent = Cm(0.74)
    bullet.paragraph_format.first_line_indent = Cm(-0.37)
    bullet.paragraph_format.line_spacing = 1.25
    bullet.paragraph_format.space_after = Pt(5)
    bullet.paragraph_format.widow_control = True


def set_run_font(run, bold: bool | None = None) -> None:
    run.font.name = LATIN_FONT
    run.font.color.rgb = BLACK
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), LATIN_FONT)
    rfonts.set(qn("w:hAnsi"), LATIN_FONT)
    rfonts.set(qn("w:eastAsia"), CN_FONT)
    rfonts.set(qn("w:cs"), LATIN_FONT)


def add_text_run(paragraph, text: str, bold: bool | None = None):
    run = paragraph.add_run(clean_text(text))
    set_run_font(run, bold)
    return run


def add_paragraph(document: Document, text: str, style_name: str):
    paragraph = document.add_paragraph(style=style_name)
    add_text_run(paragraph, text)
    return paragraph


def add_panel_spacer(document: Document) -> None:
    """Match the retained 0.2-line blank spacer between panel groups."""
    paragraph = document.add_paragraph(style=STYLE_BODY)
    paragraph.paragraph_format.line_spacing = 0.2
    paragraph.paragraph_format.first_line_indent = Cm(0)


def add_labeled_body(document: Document, label: str, text: str):
    paragraph = document.add_paragraph(style=STYLE_BODY)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.keep_together = True
    add_text_run(paragraph, label, True)
    add_text_run(paragraph, text)
    return paragraph


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style=STYLE_BULLET)
    add_text_run(paragraph, text)


def image_dimensions(path: Path, max_width_cm: float, max_height_cm: float):
    with Image.open(path) as image:
        width_px, height_px = image.size
    if width_px <= 0 or height_px <= 0:
        raise ValueError(f"Invalid image dimensions: {path}")
    aspect = width_px / height_px
    width_cm = min(max_width_cm, max_height_cm * aspect)
    height_cm = width_cm / aspect
    return width_cm, height_cm


def add_picture(
    document: Document,
    path: Path,
    alt_text: str,
    max_width_cm: float,
    max_height_cm: float,
) -> None:
    width_cm, height_cm = image_dimensions(path, max_width_cm, max_height_cm)
    paragraph = document.add_paragraph(style=STYLE_PICTURE)
    run = paragraph.add_run()
    inline_shape = run.add_picture(
        str(path), width=Cm(width_cm), height=Cm(height_cm)
    )
    inline_shape._inline.docPr.set("descr", clean_text(alt_text))


def label_font(size: int = 28):
    candidates = [
        Path("C:/Windows/Fonts/times.ttf"),
        Path("C:/Windows/Fonts/timesbd.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def make_support_sheet(
    items: list[dict[str, Any]], base_dir: Path, output: Path
) -> Path:
    resolved = []
    for index, item in enumerate(items, start=1):
        if not isinstance(item, dict) or not clean_text(item.get("path")):
            raise ValueError(f"support_images[{index}] must contain path")
        resolved.append(
            (clean_text(item.get("label")), resolve_path(item["path"], base_dir))
        )

    if len(resolved) == 1 and not resolved[0][0]:
        return resolved[0][1]

    columns = 1 if len(resolved) == 1 else 2
    cell_width, cell_height, label_height = 1100, 760, 54
    rows = (len(resolved) + columns - 1) // columns
    sheet = Image.new("RGB", (columns * cell_width, rows * cell_height), "white")
    font = label_font()

    for index, (label, path) in enumerate(resolved):
        with Image.open(path) as source:
            image = ImageOps.exif_transpose(source).convert("RGB")
        image.thumbnail((cell_width - 60, cell_height - label_height - 30))
        canvas = Image.new("RGB", (cell_width, cell_height), "white")
        draw = ImageDraw.Draw(canvas)
        if label:
            draw.text((22, 12), label, fill="black", font=font)
        x = (cell_width - image.width) // 2
        y = label_height + (cell_height - label_height - image.height) // 2
        canvas.paste(image, (x, y))
        sheet.paste(
            canvas,
            ((index % columns) * cell_width, (index // columns) * cell_height),
        )

    output.parent.mkdir(parents=True, exist_ok=True)
    sheet.save(output, format="PNG", optimize=True)
    return output


def add_named_items(
    document: Document, heading: str, items: list[dict[str, Any]]
) -> None:
    if not items:
        return
    add_paragraph(document, heading, STYLE_SECTION)
    for item in items:
        if not isinstance(item, dict):
            raise ValueError(f"{heading} items must be objects")
        add_paragraph(document, clean_text(item.get("title")), STYLE_STEP)
        add_paragraph(document, clean_text(item.get("text")), STYLE_BODY)


def evidence_label(figure_data: dict[str, Any], index: int) -> str:
    explicit = clean_text(figure_data.get("label"))
    if explicit:
        return explicit.rstrip("：:")

    title = clean_text(figure_data.get("title"))
    figure_match = re.search(r"(?:Figure|Fig\.?)\s*(\d+[A-Za-z]?)", title, re.IGNORECASE)
    if figure_match:
        return f"Fig.{figure_match.group(1)}"
    table_match = re.search(r"Table\s*(\d+[A-Za-z]?)", title, re.IGNORECASE)
    if table_match:
        return f"Table {table_match.group(1)}"
    return f"Fig.{index}"


def figure_heading(figure_data: dict[str, Any], index: int) -> str:
    """Return the evidence label plus the translated full caption title.

    New content places the translation directly in ``title``. The legacy
    translation fields remain readable and are folded into the same heading.
    """
    translated = clean_text(
        figure_data.get("translated_title")
        or figure_data.get("caption_title_cn")
    )
    title = clean_text(figure_data.get("title"))
    if not translated:
        return title

    heading_match = re.match(
        r"^((?:Figure|Fig\.?|Table)\s*\d+[A-Za-z]?)\s*[：:]",
        title,
        re.IGNORECASE,
    )
    if heading_match:
        return f"{heading_match.group(1)}：{translated}"

    label = evidence_label(figure_data, index)
    figure_match = re.fullmatch(r"Fig\.\s*(\d+[A-Za-z]?)", label, re.IGNORECASE)
    if figure_match:
        return f"Figure {figure_match.group(1)}：{translated}"
    return f"{label}：{translated}"


def translated_figure_title(figure_data: dict[str, Any]) -> str:
    """Legacy helper retained for callers that imported it directly."""
    translated = clean_text(
        figure_data.get("translated_title")
        or figure_data.get("caption_title_cn")
    )
    if translated:
        return translated

    title = clean_text(figure_data.get("title"))
    for separator in ("：", ":"):
        if separator in title:
            fallback = clean_text(title.split(separator, 1)[1])
            if fallback:
                return fallback
    return title


def normalized_logic_text(value: Any) -> str:
    text = clean_text(value)
    match = re.match(
        r"^(对应(?:Fig\.?|Figure|Table)[^。:：]*)(?:。|:|：)(.*)$",
        text,
    )
    if not match:
        return text
    label, remainder = match.groups()
    return f"{label}：{remainder}"


def save_document(document: Document, output: Path) -> None:
    try:
        document.save(output)
    except PermissionError as error:
        raise RuntimeError(
            f"Could not write {output}: the file may be open in Word. "
            "Close it and retry."
        ) from error


def build_document(
    data: dict[str, Any],
    template: Path,
    content_dir: Path,
    output: Path,
    work_dir: Path | None = None,
) -> None:
    if template.resolve() == DEFAULT_TEMPLATE.resolve():
        actual_hash = sha256(template)
        if actual_hash != EXPECTED_TEMPLATE_SHA256:
            raise ValueError(
                "Retained template hash mismatch; inspect and redistill the template"
            )

    document = Document(template)
    clear_document_body(document)
    configure_styles(document)

    section = document.sections[0]
    section.page_width = Cm(21.59)
    section.page_height = Cm(27.94)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)

    document.core_properties.title = clean_text(data["document_title"])
    add_paragraph(document, data["document_title"], STYLE_TITLE)
    add_paragraph(document, data["citation"], STYLE_METADATA)

    positioning = data["field_positioning"]
    add_paragraph(
        document,
        clean_text(positioning.get("heading") or "领域定位与本文解决的问题"),
        STYLE_SECTION,
    )
    for paragraph in positioning["paragraphs"]:
        add_paragraph(document, paragraph, STYLE_BODY)

    bullet_notes = data.get("bullet_notes") or []
    if bullet_notes:
        for note in bullet_notes:
            add_bullet(document, clean_text(note))

    add_paragraph(document, "整篇论文的图像逻辑链", STYLE_SECTION)
    for step in data["logic_chain"]:
        add_paragraph(document, clean_text(step.get("title")), STYLE_STEP)
        add_paragraph(document, normalized_logic_text(step.get("text")), STYLE_BODY)

    future_outlook = data.get("future_outlook")
    if future_outlook is None:
        # Keep existing task JSON usable while adopting the single future-outlook section.
        future_outlook = [
            *(data.get("limitations") or []),
            *(data.get("trends") or []),
        ]
    add_named_items(document, "未来展望", future_outlook)

    if work_dir is not None:
        work_dir = work_dir.resolve()
        work_dir.mkdir(parents=True, exist_ok=True)
        temp_context = nullcontext(work_dir)
    else:
        temp_context = tempfile.TemporaryDirectory(prefix="figure-evidence-")

    with temp_context as temp_dir_raw:
        temp_dir = Path(temp_dir_raw)
        evidence_blocks = data.get("evidence_blocks") or data.get("figures") or []
        for index, figure_data in enumerate(evidence_blocks, start=1):
            heading = figure_heading(figure_data, index)
            add_paragraph(document, heading, STYLE_FIGURE)
            main_image = resolve_path(figure_data["image"], content_dir)
            add_picture(
                document,
                main_image,
                f"{heading}完整图表和说明",
                15.5,
                18.3,
            )
            add_paragraph(document, figure_data["paper_claim"], STYLE_BODY)
            panel_explanations = figure_data["panel_explanations"]
            for panel_index, panel_data in enumerate(panel_explanations):
                if isinstance(panel_data, dict):
                    add_paragraph(document, panel_data["title"], STYLE_STEP)
                    add_paragraph(document, panel_data["text"], STYLE_BODY)
                else:
                    # Legacy content remains readable; new tasks use title/text objects.
                    add_paragraph(document, panel_data, STYLE_BODY)
                if panel_index < len(panel_explanations) - 1:
                    add_panel_spacer(document)
            label = evidence_label(figure_data, index)
            add_labeled_body(document, f"{label}总结：", figure_data["logic_role"])

            support_text = clean_text(figure_data.get("support_text"))
            support_images = figure_data.get("support_images") or []
            if support_text:
                add_labeled_body(document, "补充材料如何支撑：", support_text)
            if support_images:
                add_paragraph(document, "相关补充材料组图", STYLE_SUPPLEMENT)
                support_sheet = make_support_sheet(
                    support_images,
                    content_dir,
                    temp_dir / f"support-{index}.png",
                )
                add_picture(
                    document,
                    support_sheet,
                    f"{heading}相关补充材料组图",
                    14.6,
                    16.5,
                )

        output.parent.mkdir(parents=True, exist_ok=True)
        save_document(document, output)


def main() -> None:
    args = parse_args()
    content_path = args.content.resolve()
    template_path = args.template.resolve()
    output_path = args.output.resolve()

    with content_path.open("r", encoding="utf-8-sig") as stream:
        data = json.load(stream)
    if not isinstance(data, dict):
        raise ValueError("Content JSON must contain an object at the top level")
    validate_content(data)
    build_document(
        data,
        template_path,
        content_path.parent,
        output_path,
        args.work_dir.resolve() if args.work_dir else None,
    )
    print(output_path)


if __name__ == "__main__":
    main()
