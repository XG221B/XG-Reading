#!/usr/bin/env python3
"""Audit the deterministic formatting contract of a Figure Evidence DOCX."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from pathlib import Path

from build_reader_docx import (
    MATERIALS_HEADING,
    MATERIALS_SUMMARY_LABEL,
    MATERIALS_TABLE_COLUMN_WIDTHS_CM,
    MATERIALS_TABLE_HEADERS,
    clean_text,
    evidence_label,
    figure_heading,
    image_dimensions,
    materials_structure_heading,
    normalized_logic_text,
)
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

CN_FONT = "华文中宋"
LATIN_FONT = "Times New Roman"
CUSTOM_STYLES = {
    "FER Title",
    "FER Metadata",
    "FER Section",
    "FER Step",
    "FER Body",
    "FER Figure",
    "FER Figure Translation",
    "FER Supplement",
    "FER Picture",
    "FER Bullet",
}

CENTER = WD_ALIGN_PARAGRAPH.CENTER
LEFT = WD_ALIGN_PARAGRAPH.LEFT
JUSTIFY = WD_ALIGN_PARAGRAPH.JUSTIFY
CJK = r"\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff"
BOUNDARY_SPACE = re.compile(rf"(?:[{CJK}]\s+[A-Za-z0-9]|[A-Za-z0-9]\s+[{CJK}])")

STYLE_CONTRACT = {
    "FER Title": {
        "size": Pt(16), "bold": True, "alignment": CENTER,
        "line_spacing": 1.25, "space_after": Pt(8), "keep_with_next": False,
    },
    "FER Metadata": {
        "size": Pt(11), "bold": False, "alignment": CENTER,
        "line_spacing": 1.25, "space_after": Pt(12), "keep_with_next": False,
    },
    "FER Section": {
        "size": Pt(13), "bold": True, "alignment": CENTER,
        "line_spacing": 1.25, "space_before": Pt(10), "space_after": Pt(8),
        "keep_with_next": True, "widow_control": True,
    },
    "FER Step": {
        "size": Pt(11), "bold": True, "alignment": LEFT,
        "line_spacing": 1.25, "space_after": Pt(3),
        "keep_with_next": True, "widow_control": True,
    },
    "FER Body": {
        "size": Pt(11), "bold": False, "alignment": JUSTIFY,
        "line_spacing": 1.25, "space_after": Pt(6),
        "first_line_indent": Cm(0.74), "widow_control": True,
    },
    "FER Figure": {
        "size": Pt(13), "bold": True, "alignment": CENTER,
        "line_spacing": 1.25, "space_after": Pt(8),
        "keep_with_next": True, "widow_control": True,
    },
    "FER Figure Translation": {
        "size": Pt(11), "bold": True, "alignment": CENTER,
        "line_spacing": 1.25, "space_before": Pt(6), "space_after": Pt(4),
        "keep_with_next": True, "widow_control": True,
    },
    "FER Supplement": {
        "size": Pt(11), "bold": True, "alignment": CENTER,
        "line_spacing": 1.25, "space_after": Pt(2), "keep_with_next": True,
    },
    "FER Picture": {
        "size": Pt(11), "bold": False, "alignment": CENTER,
        "space_after": Pt(4),
    },
    "FER Bullet": {
        "size": Pt(11), "bold": False, "alignment": JUSTIFY,
        "line_spacing": 1.25, "space_after": Pt(5),
        "left_indent": Cm(0.74), "first_line_indent": Cm(-0.37),
        "widow_control": True,
    },
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--docx", required=True, type=Path)
    parser.add_argument("--content", type=Path)
    parser.add_argument("--json-report", type=Path)
    return parser.parse_args()


def close_enough(actual, expected, tolerance: int = 1500) -> bool:
    return actual is not None and abs(int(actual) - int(expected)) <= tolerance


def line_spacing_close_enough(actual, expected) -> bool:
    return actual is not None and abs(float(actual) - float(expected)) <= 0.01


def inherited_font(style, attr: str):
    current = style
    while current is not None:
        rpr = current._element.rPr
        if rpr is not None and rpr.rFonts is not None:
            value = rpr.rFonts.get(qn(f"w:{attr}"))
            if value:
                return value
        current = current.base_style
    return None


def inherited_size(style):
    current = style
    while current is not None:
        if current.font.size is not None:
            return current.font.size
        current = current.base_style
    return None


def emu(value) -> int:
    return int(value) if value is not None else 0


def flag_set(actual, expected) -> bool:
    if expected:
        return actual is True
    return actual in (None, False)


def check_style_contract(
    document: Document, name: str, expected: dict, errors: list[str]
) -> None:
    if name not in document.styles:
        errors.append(f"Missing style: {name}")
        return
    style = document.styles[name]
    if not close_enough(inherited_size(style), expected["size"], tolerance=10):
        errors.append(f"Incorrect size for {name}")
    if style.paragraph_format.alignment != expected["alignment"]:
        errors.append(f"Incorrect alignment for {name}")
    if inherited_font(style, "ascii") != LATIN_FONT:
        errors.append(f"Incorrect Latin font for {name}")
    if inherited_font(style, "eastAsia") != CN_FONT:
        errors.append(f"Incorrect Chinese font for {name}")
    if not flag_set(style.font.bold, expected["bold"]):
        errors.append(f"Incorrect bold for {name}")
    paragraph_format = style.paragraph_format
    if "line_spacing" in expected and not line_spacing_close_enough(
        paragraph_format.line_spacing, expected["line_spacing"]
    ):
        errors.append(f"Incorrect line spacing for {name}")
    for attribute in ("space_before", "space_after", "first_line_indent", "left_indent"):
        if attribute not in expected:
            continue
        wanted = expected[attribute]
        actual = getattr(paragraph_format, attribute)
        if not close_enough(emu(actual), emu(wanted)):
            errors.append(f"Incorrect {attribute} for {name}")
    if "keep_with_next" in expected and not flag_set(
        paragraph_format.keep_with_next, expected["keep_with_next"]
    ):
        errors.append(f"Incorrect keep_with_next for {name}")
    if "widow_control" in expected and not flag_set(
        paragraph_format.widow_control, expected["widow_control"]
    ):
        errors.append(f"Incorrect widow_control for {name}")


def is_subsequence(needle: list[str], haystack: list[str]) -> bool:
    haystack_index = 0
    for item in needle:
        while haystack_index < len(haystack) and haystack[haystack_index] != item:
            haystack_index += 1
        if haystack_index == len(haystack):
            return False
        haystack_index += 1
    return True


def build_expected_texts(data: dict) -> list[str]:
    expected = [clean_text(data["document_title"]), clean_text(data["citation"])]
    materials = data.get("materials")
    if isinstance(materials, dict) and materials.get("items"):
        expected.append(clean_text(materials.get("heading") or MATERIALS_HEADING))
        structure = materials.get("structure")
        if isinstance(structure, dict):
            expected.append(clean_text(materials_structure_heading(structure)))
            expected.append(clean_text(structure.get("explanation")))
        expected.append(
            clean_text(f"{MATERIALS_SUMMARY_LABEL}{materials.get('summary')}")
        )
    positioning = data["field_positioning"]
    expected.append(clean_text(positioning.get("heading") or "领域定位与本文解决的问题"))
    expected.extend(clean_text(item) for item in positioning["paragraphs"])
    expected.extend(clean_text(note) for note in data.get("bullet_notes") or [])
    expected.append(clean_text("整篇论文的图像逻辑链"))
    for step in data["logic_chain"]:
        expected.append(clean_text(step.get("title")))
        expected.append(normalized_logic_text(step.get("text")))
    future = data.get("future_outlook")
    if future is None:
        future = [
            *(data.get("limitations") or []),
            *(data.get("trends") or []),
        ]
    if future:
        expected.append(clean_text("未来展望"))
        for item in future:
            expected.append(clean_text(item.get("title")))
            expected.append(clean_text(item.get("text")))
    blocks = data.get("evidence_blocks") or data.get("figures") or []
    for index, block in enumerate(blocks, start=1):
        expected.append(clean_text(figure_heading(block, index)))
        expected.append(clean_text(block.get("paper_claim")))
        for panel in block["panel_explanations"]:
            if isinstance(panel, dict):
                expected.append(clean_text(panel["title"]))
                expected.append(clean_text(panel["text"]))
            else:
                expected.append(clean_text(panel))
        label = evidence_label(block, index)
        expected.append(clean_text(f"{label}总结：{block.get('logic_role')}"))
        support_text = clean_text(block.get("support_text"))
        if support_text:
            expected.append(clean_text(f"补充材料如何支撑：{support_text}"))
        if block.get("support_images"):
            expected.append(clean_text("相关补充材料组图"))
    return expected


def expected_image_widths(data: dict, content_dir: Path) -> list[tuple[str, int | None]]:
    """Expected widths per inline shape: exact for main images (recomputed with
    the same aspect limits the builder uses), a max bound for generated
    support sheets whose source images are not available to the audit."""
    widths: list[tuple[str, int | None]] = []
    blocks = data.get("evidence_blocks") or data.get("figures") or []
    for block in blocks:
        raw = clean_text(block.get("image"))
        exact = None
        if raw:
            path = Path(raw)
            if not path.is_absolute():
                path = content_dir / path
            try:
                width_cm, _ = image_dimensions(path, 15.5, 18.3)
                exact = int(Cm(width_cm))
            except (OSError, ValueError):
                exact = None
        widths.append(("exact", exact))
        if block.get("support_images"):
            widths.append(("max", int(Cm(14.6))))
    return widths


def materials_structure_image(content: dict, content_dir: Path) -> tuple[int, int | None]:
    """Return (image count, exact width EMU) for the materials structure figure."""
    materials = content.get("materials")
    if not isinstance(materials, dict):
        return 0, None
    structure = materials.get("structure")
    if not isinstance(structure, dict) or not clean_text(structure.get("image")):
        return 0, None
    path = Path(clean_text(structure["image"]))
    if not path.is_absolute():
        path = content_dir / path
    try:
        width_cm, _ = image_dimensions(path, 15.5, 18.3)
        return 1, int(Cm(width_cm))
    except (OSError, ValueError):
        return 1, None


def table_has_black_grid_borders(table) -> bool:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        return False
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            return False
        if element.get(qn("w:val")) != "single":
            return False
        if element.get(qn("w:sz")) != "6":
            return False
        color = (element.get(qn("w:color")) or "").upper()
        if color not in ("000000", "AUTO"):
            return False
    return True


def table_has_no_shading(table) -> bool:
    if table._tbl.tblPr.find(qn("w:shd")) is not None:
        return False
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.tcPr
            if tc_pr is not None and tc_pr.find(qn("w:shd")) is not None:
                return False
    return True


def check_table_cell(
    cell, expected_text, *, bold: bool, alignment, label: str, errors: list[str]
) -> None:
    raw_text = cell.text
    expected = clean_text(expected_text)
    actual = clean_text(raw_text)
    if actual != expected:
        errors.append(f"{label} text {actual!r} != expected {expected!r}")
    if BOUNDARY_SPACE.search(raw_text):
        errors.append(f"{label} contains a Chinese-Latin boundary space")
    paragraph = cell.paragraphs[0]
    if paragraph.style.name != "FER Body":
        errors.append(f"{label} paragraph style {paragraph.style.name!r} is not FER Body")
    if paragraph.paragraph_format.alignment != alignment:
        errors.append(f"{label} alignment is not the retained table alignment")
    if not close_enough(emu(paragraph.paragraph_format.first_line_indent), 0):
        errors.append(f"{label} first-line indent is not zero")
    if not line_spacing_close_enough(paragraph.paragraph_format.line_spacing, 1.15):
        errors.append(f"{label} line spacing is not 1.15")
    if not close_enough(emu(paragraph.paragraph_format.space_after), emu(Pt(2))):
        errors.append(f"{label} space after is not 2 pt")
    if not close_enough(emu(paragraph.paragraph_format.space_before), emu(Pt(2))):
        errors.append(f"{label} space before is not 2 pt")
    if cell.vertical_alignment != WD_ALIGN_VERTICAL.CENTER:
        errors.append(f"{label} is not vertically centered")
    for run in paragraph.runs:
        if not run.text.strip():
            continue
        rpr = run._element.rPr
        if rpr is None or rpr.rFonts is None:
            errors.append(f"{label} run lacks explicit fonts")
            continue
        if rpr.rFonts.get(qn("w:ascii")) != LATIN_FONT:
            errors.append(f"{label} run Latin font is wrong")
        if rpr.rFonts.get(qn("w:eastAsia")) != CN_FONT:
            errors.append(f"{label} run Chinese font is wrong")
        color = run.font.color.rgb
        if color is not None and str(color) != "000000":
            errors.append(f"{label} run is not black")
        if run.bold is not bold:
            errors.append(f"{label} run bold is {run.bold}, expected {bold}")


def check_materials(
    document: Document, content: dict, content_dir: Path, errors: list[str]
) -> None:
    materials = content.get("materials")
    tables = document.tables
    if materials is None:
        if tables:
            errors.append("Unexpected table: content.json has no materials section")
        return
    if not isinstance(materials, dict):
        errors.append("materials must be an object")
        return
    items = materials.get("items")
    if not isinstance(items, list) or not items:
        errors.append("materials.items must be a non-empty array")
        return
    if len(tables) != 1:
        errors.append(f"Materials table count {len(tables)} != expected 1")
        return
    table = tables[0]
    if table.alignment != WD_TABLE_ALIGNMENT.CENTER:
        errors.append("Materials table is not centered")
    if len(table.columns) != len(MATERIALS_TABLE_HEADERS):
        errors.append(
            f"Materials table columns {len(table.columns)} != "
            f"{len(MATERIALS_TABLE_HEADERS)}"
        )
        return
    for index, width_cm in enumerate(MATERIALS_TABLE_COLUMN_WIDTHS_CM):
        if not close_enough(emu(table.columns[index].width), emu(Cm(width_cm))):
            errors.append(f"Materials table column {index + 1} width is not {width_cm} cm")
    if len(table.rows) != len(items) + 1:
        errors.append(
            f"Materials table rows {len(table.rows)} != expected {len(items) + 1}"
        )
        return
    for row_index, row in enumerate(table.rows, start=1):
        tr_pr = row._tr.trPr
        if tr_pr is None or tr_pr.find(qn("w:cantSplit")) is None:
            errors.append(f"Materials table row {row_index} may split across pages")
    for column, header in enumerate(MATERIALS_TABLE_HEADERS):
        check_table_cell(
            table.cell(0, column),
            header,
            bold=True,
            alignment=CENTER,
            label=f"Materials header column {column + 1}",
            errors=errors,
        )
    for row_index, item in enumerate(items, start=1):
        check_table_cell(
            table.cell(row_index, 0),
            item.get("name"),
            bold=True,
            alignment=LEFT,
            label=f"Materials row {row_index} name",
            errors=errors,
        )
        check_table_cell(
            table.cell(row_index, 1),
            item.get("role"),
            bold=False,
            alignment=LEFT,
            label=f"Materials row {row_index} role",
            errors=errors,
        )
    if not table_has_black_grid_borders(table):
        errors.append("Materials table does not use the retained black single-line grid")
    if not table_has_no_shading(table):
        errors.append("Materials table must not use table or cell shading")

    paragraphs = document.paragraphs
    heading_text = clean_text(materials.get("heading") or MATERIALS_HEADING)
    heading_paragraph = paragraphs[2] if len(paragraphs) > 2 else None
    if heading_paragraph is None or clean_text(heading_paragraph.text) != heading_text:
        errors.append("Materials heading paragraph is missing or out of order")
    elif heading_paragraph.style.name != "FER Section":
        errors.append("Materials heading must use the FER Section style")

    structure = materials.get("structure")
    if isinstance(structure, dict):
        structure_text = clean_text(materials_structure_heading(structure))
        structure_paragraph = next(
            (paragraph for paragraph in paragraphs if clean_text(paragraph.text) == structure_text),
            None,
        )
        if structure_paragraph is None:
            errors.append("Materials structure heading paragraph is missing")
        else:
            if structure_paragraph.style.name != "FER Step":
                errors.append("Materials structure heading must use the FER Step style")
            if not close_enough(
                emu(structure_paragraph.paragraph_format.space_before), emu(Pt(6))
            ):
                errors.append("Materials structure heading must keep 6 pt space before")


def inline_shape_sha256(document: Document, shape) -> str:
    relationship_id = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
    blob = document.part.related_parts[relationship_id].blob
    return hashlib.sha256(blob).hexdigest()


def expected_image_manifest(content: dict) -> list[tuple[str, str | None]]:
    expected: list[tuple[str, str | None]] = []
    materials = content.get("materials")
    if isinstance(materials, dict):
        structure = materials.get("structure")
        if isinstance(structure, dict) and clean_text(structure.get("image")):
            expected.append(("materials_structure", materials_structure_heading(structure)))
    blocks = content.get("evidence_blocks") or content.get("figures") or []
    for index, block in enumerate(blocks, start=1):
        label = evidence_label(block, index)
        expected.append(("evidence_main", label))
        if block.get("support_images"):
            expected.append(("support_sheet", label))
    return expected


def check_image_identity(
    document: Document,
    content: dict,
    content_dir: Path,
    docx_path: Path,
    errors: list[str],
    warnings: list[str],
) -> None:
    """Verify the embedded image bytes at each occurrence against the build map."""
    manifest_path = content_dir / "image-manifest.json"
    if not manifest_path.exists():
        warnings.append(
            "image identity unverified: image-manifest.json is missing "
            "(legacy or externally built DOCX)"
        )
        return
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError) as error:
        errors.append(f"Image manifest cannot be read: {error}")
        return
    if not isinstance(manifest, dict) or manifest.get("schema_version") != 1:
        errors.append("Image manifest has an unsupported schema")
        return
    records = manifest.get("images")
    if not isinstance(records, list):
        errors.append("Image manifest images must be an array")
        return
    shapes = document.inline_shapes
    if len(records) != len(shapes):
        errors.append(
            f"Image manifest records {len(records)} != embedded image occurrences "
            f"{len(shapes)}"
        )
        return
    recorded_docx = str(manifest.get("docx_sha256") or "").lower()
    try:
        actual_docx = hashlib.sha256(docx_path.read_bytes()).hexdigest()
    except OSError:
        actual_docx = ""
    if recorded_docx and actual_docx and recorded_docx != actual_docx:
        warnings.append(
            "image manifest DOCX hash differs; the DOCX was modified after build"
        )
    expected_order = expected_image_manifest(content)
    actual_order = [
        (record.get("kind"), record.get("object"))
        for record in records
        if isinstance(record, dict)
    ]
    if expected_order != actual_order:
        errors.append(
            "Image manifest order does not match the content object order: "
            f"{actual_order!r} != {expected_order!r}"
        )
    for index, (record, shape) in enumerate(zip(records, shapes), start=1):
        if not isinstance(record, dict):
            errors.append(f"Image manifest record {index} must be an object")
            continue
        actual = inline_shape_sha256(document, shape)
        recorded = str(record.get("embedded_sha256") or "").lower()
        if not recorded or actual != recorded:
            errors.append(
                f"Image {index} embedded bytes do not match the recorded image identity"
            )
        source = record.get("source")
        source_hash = str(record.get("source_sha256") or "").lower()
        if source:
            source_path = Path(str(source))
            if not source_path.is_file():
                errors.append(f"Image {index} source file is missing: {source_path}")
            elif source_hash:
                actual_source = hashlib.sha256(source_path.read_bytes()).hexdigest()
                if actual_source != source_hash:
                    errors.append(
                        f"Image {index} source file changed since build: {source_path}"
                    )
        members = record.get("members")
        if isinstance(members, list):
            for member_index, member in enumerate(members, start=1):
                if not isinstance(member, dict):
                    continue
                member_path = Path(str(member.get("path") or ""))
                expected_hash = str(member.get("sha256") or "").lower()
                if not member_path.is_file():
                    errors.append(
                        f"Image {index} support member {member_index} is missing: "
                        f"{member_path}"
                    )
                elif expected_hash:
                    actual_member = hashlib.sha256(member_path.read_bytes()).hexdigest()
                    if actual_member != expected_hash:
                        errors.append(
                            f"Image {index} support member {member_index} changed since build"
                        )


def check_spacer_positions(document: Document, errors: list[str]) -> None:
    paragraphs = document.paragraphs
    meaningful = [
        index
        for index, paragraph in enumerate(paragraphs)
        if paragraph.text.strip() or paragraph._p.xpath(".//w:drawing")
    ]

    def prev_meaningful(index):
        candidates = [position for position in meaningful if position < index]
        return candidates[-1] if candidates else None

    def next_meaningful(index):
        candidates = [position for position in meaningful if position > index]
        return candidates[0] if candidates else None

    for index, paragraph in enumerate(paragraphs):
        if paragraph.text.strip() or paragraph.style.name != "FER Body":
            continue
        if paragraph.paragraph_format.line_spacing != 0.2:
            continue
        previous = prev_meaningful(index)
        following = next_meaningful(index)
        if previous is None or following is None:
            errors.append(f"Panel spacer without neighboring text at paragraph {index}")
            continue
        previous_style = paragraphs[previous].style.name
        following_style = paragraphs[following].style.name
        if previous_style != "FER Body" or following_style != "FER Step":
            errors.append(f"Panel spacer not between panel text and panel title at paragraph {index}")
            continue
        before_previous = prev_meaningful(previous)
        if before_previous is None or paragraphs[before_previous].style.name != "FER Step":
            errors.append(f"Panel spacer after a non-panel paragraph at paragraph {index}")


def audit(docx_path: Path, content_path: Path | None):
    document = Document(docx_path)
    errors: list[str] = []
    warnings: list[str] = []

    if len(document.sections) != 1:
        errors.append(f"Expected 1 section, found {len(document.sections)}")
    section = document.sections[0]
    checks = {
        "page width": (section.page_width, Cm(21.59)),
        "page height": (section.page_height, Cm(27.94)),
        "top margin": (section.top_margin, Cm(2.0)),
        "bottom margin": (section.bottom_margin, Cm(2.0)),
        "left margin": (section.left_margin, Cm(2.3)),
        "right margin": (section.right_margin, Cm(2.3)),
    }
    for label, (actual, expected) in checks.items():
        if not close_enough(actual, expected):
            errors.append(f"Incorrect {label}: {actual}, expected {expected}")

    for name in CUSTOM_STYLES:
        check_style_contract(document, name, STYLE_CONTRACT[name], errors)

    for name in CUSTOM_STYLES:
        if name == "FER Figure" or name not in document.styles:
            continue
        if document.styles[name].paragraph_format.page_break_before is True:
            errors.append(f"Unexpected page break before style {name}")

    if "FER Figure" in document.styles:
        figure_style = document.styles["FER Figure"]
        if figure_style.paragraph_format.page_break_before is not True:
            errors.append("FER Figure must start on a new page")

    cjk = r"\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff"
    boundary_space = re.compile(
        rf"(?:[{cjk}]\s+[A-Za-z0-9]|[A-Za-z0-9]\s+[{cjk}])"
    )
    figure_headings = 0
    translated_figure_headings = 0
    panel_spacers = 0
    forbidden_prompts = (
        "论文中的说法：",
        "图上如何体现：",
        "这张图在论文逻辑中的位置：",
    )
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if not text and paragraph.style.name == "FER Body":
            if paragraph.paragraph_format.line_spacing == 0.2:
                panel_spacers += 1
                if paragraph.paragraph_format.first_line_indent != Cm(0):
                    errors.append(
                        f"Panel spacer first-line indent is not zero at paragraph {index}"
                    )
            elif not paragraph._p.xpath(".//w:drawing"):
                errors.append(f"Unexpected blank FER Body paragraph at paragraph {index}")
        if not text and not paragraph._p.xpath(".//w:drawing"):
            continue
        style_name = paragraph.style.name
        if style_name.startswith("Heading") or style_name == "Title":
            errors.append(f"Collapsible built-in heading style at paragraph {index}")
        if text and style_name not in CUSTOM_STYLES:
            errors.append(
                f"Unexpected paragraph style {style_name!r} at paragraph {index}"
            )
        if style_name == "FER Figure":
            figure_headings += 1
        if style_name == "FER Figure Translation":
            translated_figure_headings += 1
        if any(text.startswith(prompt) for prompt in forbidden_prompts):
            errors.append(f"Deprecated prompt at paragraph {index}: {text}")
        if boundary_space.search(text):
            errors.append(f"Chinese-Latin boundary space at paragraph {index}")

        for run in paragraph.runs:
            if not run.text.strip():
                continue
            rpr = run._element.rPr
            if rpr is None or rpr.rFonts is None:
                errors.append(f"Missing explicit run fonts at paragraph {index}")
                continue
            if rpr.rFonts.get(qn("w:ascii")) != LATIN_FONT:
                errors.append(f"Wrong run Latin font at paragraph {index}")
            if rpr.rFonts.get(qn("w:eastAsia")) != CN_FONT:
                errors.append(f"Wrong run Chinese font at paragraph {index}")
            color = run.font.color.rgb
            if color is not None and str(color) != "000000":
                errors.append(f"Non-black text at paragraph {index}")

    check_spacer_positions(document, errors)

    for section in document.sections:
        for part_name, part in (("header", section.header), ("footer", section.footer)):
            if any(paragraph.text.strip() for paragraph in part.paragraphs):
                errors.append(f"Non-empty {part_name} text is not allowed")

    for index, paragraph in enumerate(document.paragraphs):
        text_runs = [run for run in paragraph.runs if run.text.strip()]
        if len(text_runs) >= 2:
            if text_runs[0].bold is not True:
                errors.append(f"Multi-run paragraph label not bold at paragraph {index}")
            if any(run.bold is True for run in text_runs[1:]):
                errors.append(f"Multi-run paragraph body run bold at paragraph {index}")

    if content_path:
        with content_path.open("r", encoding="utf-8-sig") as stream:
            content = json.load(stream)
        check_materials(document, content, content_path.parent, errors)
        check_image_identity(
            document, content, content_path.parent, docx_path, errors, warnings
        )
        if content.get("materials") is None:
            warnings.append(
                "materials section absent: required for papers centered on an "
                "engineered material, device, or micro/nanorobotic system"
            )
        if "future_outlook" not in content and (
            content.get("limitations") or content.get("trends")
        ):
            warnings.append(
                "legacy limitations/trends are rendered as 未来展望 but are not "
                "verified against the authors' Discussion/Conclusion"
            )
        evidence_blocks = content.get("evidence_blocks") or content.get("figures") or []
        expected_figures = len(evidence_blocks)
        if figure_headings != expected_figures:
            errors.append(
                f"Figure heading count {figure_headings} != expected {expected_figures}"
            )
        if translated_figure_headings:
            errors.append(
                "Separate translated figure-title paragraphs are no longer allowed"
            )
        expected_panel_spacers = sum(
            max(len(figure.get("panel_explanations") or []) - 1, 0)
            for figure in evidence_blocks
        )
        if panel_spacers != expected_panel_spacers:
            errors.append(
                f"Panel spacer count {panel_spacers} != expected {expected_panel_spacers}"
            )
        structure_images, structure_width = materials_structure_image(
            content, content_path.parent
        )
        expected_images = (
            structure_images
            + expected_figures
            + sum(1 for figure in evidence_blocks if figure.get("support_images"))
        )
        if len(document.inline_shapes) != expected_images:
            errors.append(
                f"Inline image count {len(document.inline_shapes)} != expected {expected_images}"
            )
        widths: list[tuple[str, int | None]] = []
        if structure_images:
            widths.append(("exact", structure_width))
        widths.extend(expected_image_widths(content, content_path.parent))
        if len(document.inline_shapes) == len(widths):
            for shape_index, shape in enumerate(document.inline_shapes):
                kind, expected_width = widths[shape_index]
                if expected_width is None:
                    continue
                if kind == "exact":
                    if not close_enough(shape.width, expected_width, tolerance=20000):
                        errors.append(
                            f"Inline image {shape_index + 1} width "
                            f"{shape.width} != expected {expected_width}"
                        )
                elif shape.width > expected_width + 20000:
                    errors.append(
                        f"Inline image {shape_index + 1} width "
                        f"{shape.width} exceeds support-sheet bound {expected_width}"
                    )
        doc_texts = [
            clean_text(paragraph.text)
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]
        expected_texts = build_expected_texts(content)
        if not is_subsequence(expected_texts, doc_texts):
            errors.append(
                "DOCX text is missing or altered relative to content.json"
            )
        if not is_subsequence(doc_texts, expected_texts):
            errors.append(
                "DOCX contains text not present in content.json"
            )

    if not Path("C:/Windows/Fonts/STZHONGS.TTF").exists():
        warnings.append("华文中宋字体文件未在Windows标准字体目录中找到")

    return {
        "docx": str(docx_path.resolve()),
        "errors": sorted(set(errors)),
        "warnings": sorted(set(warnings)),
        "paragraphs": len(document.paragraphs),
        "images": len(document.inline_shapes),
        "figure_headings": figure_headings,
        "translated_figure_headings": translated_figure_headings,
        "panel_spacers": panel_spacers,
    }


def main() -> None:
    args = parse_args()
    report = audit(args.docx.resolve(), args.content.resolve() if args.content else None)
    rendered = json.dumps(report, ensure_ascii=False, indent=2)
    print(json.dumps(report, ensure_ascii=True, indent=2))
    if args.json_report:
        args.json_report.parent.mkdir(parents=True, exist_ok=True)
        args.json_report.write_text(rendered, encoding="utf-8")
    if report["errors"]:
        sys.exit(1)


if __name__ == "__main__":
    main()
