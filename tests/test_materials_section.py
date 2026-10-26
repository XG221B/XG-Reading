"""Materials-system section acceptance tests (材料体系与设计思路).

Contract under test:

* Build order: title -> citation -> materials section (heading, table, optional
  structure figure, summary) -> field positioning.
* Table: 2 columns, retained header/row text, bold rules, explicit fonts, black
  text, black single-line grid, centered, fixed 4.4 + 12.59 cm columns, and rows
  that cannot split across pages.
* Structure figure: ``FER Step`` heading (``Fig.1A:<title>`` or the bare title
  when ``label`` is omitted), centered original image within 15.5 x 18.3 cm, and
  a ``FER Body`` explanation after it.
* Summary: bold ``设计思路:`` label followed by body text in a ``FER Body``
  paragraph.
* Audit: a valid materials document passes; tampering is reported; legacy
  content without materials only warns; a table without a materials section in
  content.json is an error.
* Preflight: materials structure errors, unresolvable structure image, and a
  missing materials section as a warning only.
"""

from __future__ import annotations

import json
import re
from copy import deepcopy
from pathlib import Path

import audit_reader_docx as audit
import build_reader_docx as builder
import preflight_content as preflight
import pytest
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

MATERIALS_HEADING = "材料体系与设计思路"
SUMMARY_LABEL = "设计思路："
STRUCTURE_TITLE = "材料结合后的结构示意"
EXPLANATION = "各组分通过静电吸附结合为核壳结构，为响应性释放提供结构基础。"
SUMMARY = "本文用核壳结构材料体系以响应性释放的方式解决了药物在肿瘤部位富集不足的问题。"
ITEMS = [
    {"name": "介孔二氧化硅纳米颗粒", "role": "负载药物并在肿瘤微环境中响应释放"},
    {"name": "透明质酸外壳", "role": "提高血液循环稳定性并靶向肿瘤细胞"},
]
POSITIONING_HEADING = "领域定位与本文解决的问题"


def make_content(
    tmp_path,
    *,
    with_materials=True,
    structure=True,
    structure_label="Fig.1A",
    structure_size=(400, 300),
):
    """Build a valid content.json dict with a tiny evidence block."""
    Image.new("RGB", (80, 80), "white").save(tmp_path / "fig1.png")
    if structure:
        Image.new("RGB", structure_size, "white").save(tmp_path / "structure.png")
    content = {
        "document_title": "题目与逐图解读",
        "citation": "论文：Author et al., Title, Journal,Year",
        "field_positioning": {
            "heading": POSITIONING_HEADING,
            "paragraphs": ["领域问题", "具体缺口", "本文贡献与边界"],
        },
        "logic_chain": [
            {
                "title": "Step1：为什么需要这一步",
                "text": "对应Fig.1：实验、结果及其支持的结论。",
            }
        ],
        "evidence_blocks": [
            {
                "kind": "figure",
                "label": "Fig.1",
                "title": "Figure 1：主图英文标题的准确中文翻译",
                "image": "fig1.png",
                "paper_claim": "作者对该证据块的主要表述。",
                "panel_explanations": [
                    {
                        "title": "Fig.1A：该panel图注内容的中文译名",
                        "text": "Fig.1A的读图方法与意义。",
                    }
                ],
                "logic_role": "该证据块在全文论证中的作用。",
            }
        ],
    }
    if with_materials:
        materials = {"items": deepcopy(ITEMS), "summary": SUMMARY}
        if structure:
            structure_block = {
                "title": STRUCTURE_TITLE,
                "image": "structure.png",
                "explanation": EXPLANATION,
            }
            if structure_label is not None:
                structure_block["label"] = structure_label
            materials["structure"] = structure_block
        content["materials"] = materials
    return content


def write_content(tmp_path, content, name="content.json"):
    path = tmp_path / name
    path.write_text(json.dumps(content, ensure_ascii=False), encoding="utf-8")
    return path


def build_docx(tmp_path, content, name="reader.docx"):
    path = tmp_path / name
    builder.build_document(content, builder.DEFAULT_TEMPLATE, tmp_path, path)
    return path


def make_built(tmp_path, **kwargs):
    content = make_content(tmp_path, **kwargs)
    content_path = write_content(tmp_path, content)
    docx_path = build_docx(tmp_path, content)
    return content, content_path, docx_path


@pytest.fixture
def built(tmp_path):
    return make_built(tmp_path)


def find_paragraph(document, text):
    return next(paragraph for paragraph in document.paragraphs if paragraph.text == text)


def body_sequence(document):
    """Return body children in document order as ('paragraph'|'table', object)."""
    paragraphs = iter(document.paragraphs)
    tables = iter(document.tables)
    sequence = []
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            sequence.append(("paragraph", next(paragraphs)))
        elif child.tag == qn("w:tbl"):
            sequence.append(("table", next(tables)))
    return sequence


def sequence_index(sequence, predicate):
    return next(
        index
        for index, (kind, item) in enumerate(sequence)
        if predicate(kind, item)
    )


def cell_run(cell):
    return cell.paragraphs[0].runs[0]


class TestBuildMaterialsSection:
    def test_materials_section_precedes_field_positioning(self, built):
        content, _, docx_path = built
        document = Document(docx_path)
        sequence = body_sequence(document)

        def paragraph_index(predicate):
            return sequence_index(
                sequence,
                lambda kind, item: kind == "paragraph" and predicate(item),
            )

        structure_heading = f"Fig.1A：{STRUCTURE_TITLE}"
        positions = [
            paragraph_index(lambda p: p.text == content["document_title"]),
            paragraph_index(lambda p: p.text == content["citation"]),
            paragraph_index(lambda p: p.text == MATERIALS_HEADING),
            sequence_index(sequence, lambda kind, _: kind == "table"),
            paragraph_index(lambda p: p.text == structure_heading),
            paragraph_index(
                lambda p: p.style.name == "FER Picture" and not p.text
            ),
            paragraph_index(lambda p: p.text == EXPLANATION),
            paragraph_index(lambda p: p.text == SUMMARY_LABEL + SUMMARY),
            paragraph_index(lambda p: p.text == POSITIONING_HEADING),
        ]
        assert positions == sorted(positions), positions
        assert positions[2] < positions[3] < positions[4]

    def test_materials_section_uses_retained_styles(self, built):
        _, _, docx_path = built
        document = Document(docx_path)
        heading = find_paragraph(document, MATERIALS_HEADING)
        assert heading.style.name == "FER Section"
        assert (
            document.styles["FER Section"].paragraph_format.alignment
            == WD_ALIGN_PARAGRAPH.CENTER
        )
        structure_heading = find_paragraph(document, f"Fig.1A：{STRUCTURE_TITLE}")
        assert structure_heading.style.name == "FER Step"
        assert structure_heading.paragraph_format.space_before == Pt(6)
        assert (
            document.styles["FER Step"].paragraph_format.alignment
            == WD_ALIGN_PARAGRAPH.LEFT
        )
        explanation = find_paragraph(document, EXPLANATION)
        assert explanation.style.name == "FER Body"
        assert (
            document.styles["FER Picture"].paragraph_format.alignment
            == WD_ALIGN_PARAGRAPH.CENTER
        )

    def test_table_layout_and_content(self, built):
        _, _, docx_path = built
        document = Document(docx_path)
        assert len(document.tables) == 1
        table = document.tables[0]
        assert table.alignment == WD_TABLE_ALIGNMENT.CENTER
        assert len(table.columns) == 2
        assert len(table.rows) == len(ITEMS) + 1
        assert [table.cell(0, column).text for column in range(2)] == [
            "材料/组分",
            "作用",
        ]
        for row_index, item in enumerate(ITEMS, start=1):
            assert table.cell(row_index, 0).text == item["name"]
            assert table.cell(row_index, 1).text == item["role"]

    def test_table_column_widths_fill_the_text_column(self, built):
        _, _, docx_path = built
        table = Document(docx_path).tables[0]
        widths = [table.columns[index].width for index in range(2)]
        assert abs(widths[0] - Cm(4.4)) <= 1000
        assert abs(widths[1] - Cm(12.59)) <= 1000
        assert abs(sum(widths) - Cm(16.99)) <= 2000
        assert table.autofit is False

    def test_table_cell_formatting(self, built):
        _, _, docx_path = built
        table = Document(docx_path).tables[0]
        expected_bold = {(0, 0): True, (0, 1): True}
        for row_index in range(1, len(ITEMS) + 1):
            expected_bold[(row_index, 0)] = True
            expected_bold[(row_index, 1)] = False
        for (row_index, column), bold in expected_bold.items():
            cell = table.cell(row_index, column)
            paragraph = cell.paragraphs[0]
            run = cell_run(cell)
            assert paragraph.style.name == "FER Body"
            assert cell.vertical_alignment == WD_ALIGN_VERTICAL.CENTER
            expected_alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if row_index == 0
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            assert paragraph.paragraph_format.alignment == expected_alignment
            assert paragraph.paragraph_format.first_line_indent == Cm(0)
            assert paragraph.paragraph_format.line_spacing == 1.15
            assert paragraph.paragraph_format.space_after == Pt(2)
            assert run.bold is bold
            rpr = run._element.rPr
            assert rpr is not None and rpr.rFonts is not None
            assert rpr.rFonts.get(qn("w:ascii")) == "Times New Roman"
            assert rpr.rFonts.get(qn("w:hAnsi")) == "Times New Roman"
            assert rpr.rFonts.get(qn("w:eastAsia")) == "华文中宋"
            assert str(run.font.color.rgb) == "000000"
            assert paragraph.style.font.size == Pt(11)

    def test_table_has_black_single_line_grid_without_shading(self, built):
        _, _, docx_path = built
        table = Document(docx_path).tables[0]
        borders = table._tbl.tblPr.find(qn("w:tblBorders"))
        assert borders is not None
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            element = borders.find(qn(f"w:{edge}"))
            assert element is not None, edge
            assert element.get(qn("w:val")) == "single"
            assert element.get(qn("w:sz")) == "6"
            assert element.get(qn("w:color")) == "000000"
        assert not table._tbl.xpath(".//w:shd")

    def test_table_rows_cannot_split_across_pages(self, built):
        _, _, docx_path = built
        table = Document(docx_path).tables[0]
        for row in table.rows:
            tr_pr = row._tr.trPr
            assert tr_pr is not None
            assert tr_pr.find(qn("w:cantSplit")) is not None

    def test_structure_figure_follows_heading_before_explanation(self, built):
        _, _, docx_path = built
        document = Document(docx_path)
        structure_image = Path(docx_path).parent / "structure.png"
        expected_width, expected_height = builder.image_dimensions(
            structure_image, 15.5, 18.3
        )
        assert len(document.inline_shapes) == 2  # structure + one evidence block
        shape = document.inline_shapes[0]
        assert abs(shape.width - int(Cm(expected_width))) <= 1000
        assert abs(shape.height - int(Cm(expected_height))) <= 1000
        assert shape.width <= Cm(15.5)
        assert shape.height <= Cm(18.3)
        assert shape._inline.docPr.get("descr") == f"Fig.1A：{STRUCTURE_TITLE}（原图）"

    @pytest.mark.parametrize(
        ("label", "expected_heading"),
        [
            (None, STRUCTURE_TITLE),
            ("Fig.1A", f"Fig.1A：{STRUCTURE_TITLE}"),
            ("Fig.1A：", f"Fig.1A：{STRUCTURE_TITLE}"),
        ],
    )
    def test_structure_heading_label_is_optional(
        self, tmp_path, label, expected_heading
    ):
        _, _, docx_path = make_built(tmp_path, structure_label=label)
        document = Document(docx_path)
        headings = [
            paragraph
            for paragraph in document.paragraphs
            if paragraph.style.name == "FER Step" and STRUCTURE_TITLE in paragraph.text
        ]
        assert len(headings) == 1
        assert headings[0].text == expected_heading

    @pytest.mark.parametrize("size", [(400, 300), (300, 600), (1800, 400)])
    def test_structure_image_respects_the_retained_limits(self, tmp_path, size):
        _, _, docx_path = make_built(tmp_path, structure_size=size)
        document = Document(docx_path)
        structure_image = tmp_path / "structure.png"
        expected_width, expected_height = builder.image_dimensions(
            structure_image, 15.5, 18.3
        )
        shape = document.inline_shapes[0]
        assert abs(shape.width - int(Cm(expected_width))) <= 1000
        assert abs(shape.height - int(Cm(expected_height))) <= 1000
        assert shape.width <= Cm(15.5)
        assert shape.height <= Cm(18.3)

    def test_summary_uses_a_bold_label_run(self, built):
        _, _, docx_path = built
        document = Document(docx_path)
        summary = find_paragraph(document, SUMMARY_LABEL + SUMMARY)
        assert summary.style.name == "FER Body"
        assert len(summary.runs) == 2
        assert summary.runs[0].text == SUMMARY_LABEL
        assert summary.runs[0].bold is True
        assert summary.runs[1].text == SUMMARY
        assert summary.runs[1].bold is not True
        indent = document.styles["FER Body"].paragraph_format.first_line_indent
        assert abs(indent - Cm(0.74)) <= 1000

    def test_custom_materials_heading_is_used(self, tmp_path):
        content = make_content(tmp_path)
        content["materials"]["heading"] = "本文材料体系"
        docx_path = build_docx(tmp_path, content)
        document = Document(docx_path)
        heading = find_paragraph(document, "本文材料体系")
        assert heading.style.name == "FER Section"

    def test_structure_is_optional(self, tmp_path):
        _, _, docx_path = make_built(tmp_path, structure=False)
        document = Document(docx_path)
        assert len(document.tables) == 1
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        assert MATERIALS_HEADING in paragraphs
        assert SUMMARY_LABEL + SUMMARY in paragraphs
        assert STRUCTURE_TITLE not in paragraphs
        assert len(document.inline_shapes) == 1  # only the evidence block image

    def test_legacy_content_without_materials_still_builds(self, tmp_path):
        content = make_content(tmp_path, with_materials=False)
        content.pop("materials", None)
        docx_path = build_docx(tmp_path, content)
        document = Document(docx_path)
        assert document.tables == []
        sequence = body_sequence(document)
        assert sequence[0][0] == "paragraph"
        assert sequence[0][1].text == content["document_title"]
        assert sequence[1][1].text == content["citation"]
        assert sequence[2][1].text == POSITIONING_HEADING
        assert MATERIALS_HEADING not in [
            paragraph.text for paragraph in document.paragraphs
        ]

    @pytest.mark.parametrize(
        ("case", "expected"),
        [
            ("materials_not_object", "materials must be an object"),
            ("items_missing", "materials.items must be a non-empty array"),
            ("items_empty", "materials.items must be a non-empty array"),
            ("items_not_list", "materials.items must be a non-empty array"),
            ("item_missing_name", "materials.items[2].name is required"),
            ("item_missing_role", "materials.items[2].role is required"),
            ("item_not_object", "materials.items[2] must be an object"),
            ("summary_missing", "materials.summary is required"),
            ("summary_blank", "materials.summary is required"),
            ("structure_not_object", "materials.structure must be an object"),
            ("structure_missing_title", "materials.structure.title is required"),
            ("structure_missing_image", "materials.structure.image is required"),
            (
                "structure_missing_explanation",
                "materials.structure.explanation is required",
            ),
            ("heading_blank", "materials.heading must be non-empty when present"),
        ],
    )
    def test_validate_content_rejects_bad_materials(self, tmp_path, case, expected):
        content = make_content(tmp_path)
        _break_materials(content, case)
        with pytest.raises(ValueError, match=re.escape(expected)):
            builder.validate_content(content)


def _break_materials(content, case):
    materials = content["materials"]
    if case == "materials_not_object":
        content["materials"] = "bad"
    elif case == "items_missing":
        materials.pop("items")
    elif case == "items_empty":
        materials["items"] = []
    elif case == "items_not_list":
        materials["items"] = "bad"
    elif case == "item_missing_name":
        materials["items"][1].pop("name")
    elif case == "item_missing_role":
        materials["items"][1].pop("role")
    elif case == "item_not_object":
        materials["items"][1] = "bad"
    elif case == "summary_missing":
        materials.pop("summary")
    elif case == "summary_blank":
        materials["summary"] = "   "
    elif case == "structure_not_object":
        materials["structure"] = "bad"
    elif case == "structure_missing_title":
        materials["structure"].pop("title")
    elif case == "structure_missing_image":
        materials["structure"].pop("image")
    elif case == "structure_missing_explanation":
        materials["structure"].pop("explanation")
    elif case == "heading_blank":
        materials["heading"] = "   "
    else:  # pragma: no cover - guards the parametrization
        raise AssertionError(f"unknown materials case: {case}")


class TestAuditMaterialsSection:
    def test_valid_materials_document_passes_audit(self, built):
        _, content_path, docx_path = built
        report = audit.audit(docx_path, content_path)
        assert report["errors"] == []
        assert not any(
            "materials section absent" in warning for warning in report["warnings"]
        )

    def test_tampered_table_text_is_flagged(self, built):
        _, content_path, docx_path = built
        document = Document(docx_path)
        document.tables[0].cell(1, 0).paragraphs[0].runs[0].text = "篡改材料名"
        document.save(docx_path)
        report = audit.audit(docx_path, content_path)
        assert any(
            "Materials row 1 name" in error and "!=" in error
            for error in report["errors"]
        )

    def test_removed_table_row_is_flagged(self, built):
        _, content_path, docx_path = built
        document = Document(docx_path)
        table = document.tables[0]
        table._tbl.remove(table.rows[-1]._tr)
        document.save(docx_path)
        report = audit.audit(docx_path, content_path)
        assert any("Materials table rows" in error for error in report["errors"])

    def test_unbolded_header_is_flagged(self, built):
        _, content_path, docx_path = built
        document = Document(docx_path)
        cell_run(document.tables[0].cell(0, 0)).bold = False
        document.save(docx_path)
        report = audit.audit(docx_path, content_path)
        assert any("bold" in error for error in report["errors"])

    def test_bolded_role_column_is_flagged(self, built):
        _, content_path, docx_path = built
        document = Document(docx_path)
        cell_run(document.tables[0].cell(1, 1)).bold = True
        document.save(docx_path)
        report = audit.audit(docx_path, content_path)
        assert any("Materials row 1 role" in error and "bold" in error
                   for error in report["errors"])

    def test_missing_table_borders_are_flagged(self, built):
        _, content_path, docx_path = built
        document = Document(docx_path)
        table = document.tables[0]
        table._tbl.tblPr.remove(table._tbl.tblPr.find(qn("w:tblBorders")))
        document.save(docx_path)
        report = audit.audit(docx_path, content_path)
        assert any(
            "retained black single-line grid" in error for error in report["errors"]
        )

    def test_wrong_cell_alignment_is_flagged(self, built):
        _, content_path, docx_path = built
        document = Document(docx_path)
        cell = document.tables[0].cell(0, 0)
        cell.paragraphs[0].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        document.save(docx_path)
        report = audit.audit(docx_path, content_path)
        assert any("Materials header column 1 alignment" in error
                   for error in report["errors"])

    def test_wrong_cell_font_is_flagged(self, built):
        _, content_path, docx_path = built
        document = Document(docx_path)
        run = cell_run(document.tables[0].cell(1, 0))
        run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        document.save(docx_path)
        report = audit.audit(docx_path, content_path)
        assert any("Latin font is wrong" in error for error in report["errors"])

    def test_non_black_cell_text_is_flagged(self, built):
        _, content_path, docx_path = built
        document = Document(docx_path)
        run = cell_run(document.tables[0].cell(1, 1))
        run.font.color.rgb = RGBColor(0x11, 0, 0)
        document.save(docx_path)
        report = audit.audit(docx_path, content_path)
        assert any("not black" in error for error in report["errors"])

    def test_splittable_table_row_is_flagged(self, built):
        _, content_path, docx_path = built
        document = Document(docx_path)
        row = document.tables[0].rows[1]
        row._tr.trPr.remove(row._tr.trPr.find(qn("w:cantSplit")))
        document.save(docx_path)
        report = audit.audit(docx_path, content_path)
        assert any("may split" in error for error in report["errors"])

    def test_wrong_table_column_width_is_flagged(self, built):
        _, content_path, docx_path = built
        document = Document(docx_path)
        document.tables[0].columns[0].width = Cm(3.0)
        document.save(docx_path)
        report = audit.audit(docx_path, content_path)
        assert any("Materials table column 1 width" in error
                   for error in report["errors"])

    def test_extra_image_is_flagged(self, built):
        _, content_path, docx_path = built
        document = Document(docx_path)
        document.add_picture(
            str(content_path.parent / "fig1.png"), width=Cm(2)
        )
        document.save(docx_path)
        report = audit.audit(docx_path, content_path)
        assert any("Inline image count" in error for error in report["errors"])

    def test_wrong_structure_image_width_is_flagged(self, built):
        _, content_path, docx_path = built
        document = Document(docx_path)
        document.inline_shapes[0].width = Cm(8)
        document.save(docx_path)
        report = audit.audit(docx_path, content_path)
        assert any("Inline image 1 width" in error for error in report["errors"])

    def test_missing_table_with_materials_content_is_flagged(self, tmp_path):
        materials_content = make_content(tmp_path)
        materials_path = write_content(tmp_path, materials_content, "materials.json")
        legacy_content = make_content(tmp_path, with_materials=False)
        docx_path = build_docx(tmp_path, legacy_content, "legacy.docx")
        report = audit.audit(docx_path, materials_path)
        assert any("Materials table count" in error for error in report["errors"])

    def test_table_without_materials_section_is_an_error(self, tmp_path):
        _, _, docx_path = make_built(tmp_path)
        legacy_path = write_content(
            tmp_path, make_content(tmp_path, with_materials=False), "legacy.json"
        )
        report = audit.audit(docx_path, legacy_path)
        assert any("Unexpected table" in error for error in report["errors"])

    def test_legacy_document_without_materials_only_warns(self, tmp_path):
        content = make_content(tmp_path, with_materials=False)
        content.pop("materials", None)
        content_path = write_content(tmp_path, content, "legacy.json")
        docx_path = build_docx(tmp_path, content, "legacy.docx")
        report = audit.audit(docx_path, content_path)
        assert report["errors"] == []
        assert any(
            "materials section absent" in warning for warning in report["warnings"]
        )


class TestPreflightMaterialsSection:
    def test_valid_materials_content_passes(self, tmp_path):
        result = preflight.check_content(make_content(tmp_path), tmp_path)
        assert result["errors"] == []
        assert result["warnings"] == []

    def test_structure_is_optional(self, tmp_path):
        content = make_content(tmp_path, structure=False)
        result = preflight.check_content(content, tmp_path)
        assert result["errors"] == []

    def test_absent_materials_is_only_a_warning(self, tmp_path):
        content = make_content(tmp_path, with_materials=False)
        content.pop("materials", None)
        result = preflight.check_content(content, tmp_path)
        assert result["errors"] == []
        assert any(
            "materials section absent" in warning for warning in result["warnings"]
        )

    @pytest.mark.parametrize(
        ("case", "expected"),
        [
            ("materials_not_object", "materials must be an object"),
            ("items_missing", "materials.items must be a non-empty array"),
            ("items_empty", "materials.items must be a non-empty array"),
            ("items_not_list", "materials.items must be a non-empty array"),
            ("item_missing_name", "materials.items[2].name is required"),
            ("item_missing_role", "materials.items[2].role is required"),
            ("item_not_object", "materials.items[2] must be an object"),
            ("summary_missing", "materials.summary is required"),
            ("summary_blank", "materials.summary is required"),
            ("structure_not_object", "materials.structure must be an object"),
            ("structure_missing_title", "materials.structure.title is required"),
            ("structure_missing_image", "materials.structure.image is required"),
            (
                "structure_missing_explanation",
                "materials.structure.explanation is required",
            ),
            ("heading_blank", "materials.heading must be non-empty when present"),
        ],
    )
    def test_bad_materials_are_errors(self, tmp_path, case, expected):
        content = make_content(tmp_path)
        if case == "heading_blank":
            content["materials"]["heading"] = "   "
        else:
            _break_materials(content, case)
        result = preflight.check_content(content, tmp_path)
        assert any(expected in error for error in result["errors"]), result["errors"]
        assert result["errors"]

    def test_unresolvable_structure_image_is_an_error(self, tmp_path):
        content = make_content(tmp_path)
        content["materials"]["structure"]["image"] = "nope.png"
        result = preflight.check_content(content, tmp_path)
        assert any(
            "materials.structure.image cannot be resolved: nope.png" in error
            for error in result["errors"]
        )

    def test_boundary_space_inside_materials_is_detected(self, tmp_path):
        content = make_content(tmp_path)
        content["materials"]["items"][0]["name"] = "材料 A"
        result = preflight.check_content(content, tmp_path)
        assert any("边界" in error for error in result["errors"])

    def test_forbidden_prompt_inside_materials_is_detected(self, tmp_path):
        content = make_content(tmp_path)
        content["materials"]["summary"] = "论文中的说法：材料体系良好。"
        result = preflight.check_content(content, tmp_path)
        assert any("Forbidden prompt" in error for error in result["errors"])

    def test_cli_passes_with_materials(self, tmp_path, capsys):
        content_path = write_content(tmp_path, make_content(tmp_path))
        assert preflight.main_with_args(["--content", str(content_path)]) == 0
        report = json.loads(capsys.readouterr().out)
        assert report["ok"] is True
        assert report["errors"] == []

    def test_cli_without_materials_exits_zero_with_warning(self, tmp_path, capsys):
        content = make_content(tmp_path, with_materials=False)
        content.pop("materials", None)
        content_path = write_content(tmp_path, content)
        assert preflight.main_with_args(["--content", str(content_path)]) == 0
        report = json.loads(capsys.readouterr().out)
        assert report["errors"] == []
        assert any(
            "materials section absent" in warning for warning in report["warnings"]
        )


class TestAuditHardening:
    """Regression tests for the materials-audit hardening (task-9)."""

    def test_border_line_width_tamper_is_flagged(self, built):
        _, content_path, docx_path = built
        document = Document(docx_path)
        borders = document.tables[0]._tbl.tblPr.find(qn("w:tblBorders"))
        for edge in borders:
            edge.set(qn("w:sz"), "8")
        document.save(docx_path)
        report = audit.audit(docx_path, content_path)
        assert any(
            "retained black single-line grid" in error for error in report["errors"]
        )

    def test_table_level_shading_is_flagged(self, built):
        _, content_path, docx_path = built
        document = Document(docx_path)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:fill"), "FFFF00")
        document.tables[0]._tbl.tblPr.append(shading)
        document.save(docx_path)
        report = audit.audit(docx_path, content_path)
        assert any("shading" in error for error in report["errors"])

    def test_cell_level_shading_is_flagged(self, built):
        _, content_path, docx_path = built
        document = Document(docx_path)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:fill"), "FFFF00")
        cell = document.tables[0].cell(1, 0)
        cell._tc.get_or_add_tcPr().append(shading)
        document.save(docx_path)
        report = audit.audit(docx_path, content_path)
        assert any("shading" in error for error in report["errors"])

    def test_materials_heading_wrong_style_is_flagged(self, built):
        _, content_path, docx_path = built
        document = Document(docx_path)
        find_paragraph(document, MATERIALS_HEADING).style = "FER Body"
        document.save(docx_path)
        report = audit.audit(docx_path, content_path)
        assert any(
            "Materials heading must use the FER Section style" in error
            for error in report["errors"]
        )

    def test_materials_heading_out_of_order_is_flagged(self, built):
        _, content_path, docx_path = built
        document = Document(docx_path)
        element = find_paragraph(document, MATERIALS_HEADING)._element
        element.getparent().remove(element)
        document.paragraphs[-1]._element.addnext(element)
        document.save(docx_path)
        report = audit.audit(docx_path, content_path)
        assert any(
            "Materials heading paragraph is missing or out of order" in error
            for error in report["errors"]
        )

    def test_structure_heading_wrong_style_is_flagged(self, built):
        _, content_path, docx_path = built
        document = Document(docx_path)
        find_paragraph(document, f"Fig.1A：{STRUCTURE_TITLE}").style = "FER Body"
        document.save(docx_path)
        report = audit.audit(docx_path, content_path)
        assert any(
            "Materials structure heading must use the FER Step style" in error
            for error in report["errors"]
        )

    def test_structure_heading_missing_space_before_is_flagged(self, built):
        _, content_path, docx_path = built
        document = Document(docx_path)
        paragraph = find_paragraph(document, f"Fig.1A：{STRUCTURE_TITLE}")
        paragraph.paragraph_format.space_before = Pt(0)
        document.save(docx_path)
        report = audit.audit(docx_path, content_path)
        assert any(
            "Materials structure heading must keep 6 pt space before" in error
            for error in report["errors"]
        )

    def test_missing_structure_heading_is_flagged(self, built):
        _, content_path, docx_path = built
        document = Document(docx_path)
        paragraph = find_paragraph(document, f"Fig.1A：{STRUCTURE_TITLE}")
        paragraph._element.getparent().remove(paragraph._element)
        document.save(docx_path)
        report = audit.audit(docx_path, content_path)
        assert any(
            "Materials structure heading paragraph is missing" in error
            for error in report["errors"]
        )

    def test_cell_boundary_space_is_flagged(self, tmp_path):
        content = make_content(tmp_path)
        content["materials"]["items"][0]["role"] = "pH响应释放"
        content_path = write_content(tmp_path, content)
        docx_path = build_docx(tmp_path, content)
        document = Document(docx_path)
        cell_run(document.tables[0].cell(1, 1)).text = "pH 响应释放"
        document.save(docx_path)
        report = audit.audit(docx_path, content_path)
        assert any("boundary space" in error for error in report["errors"])
