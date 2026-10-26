"""Regression tests for A07: audit image identity, not just image dimensions.

The builder records every image occurrence (order, object, source hash, support
members, embedded bytes, dimensions) in ``process/image-manifest.json``. The
auditor verifies the embedded bytes and source files page-by-page so two
same-sized figures cannot be silently swapped.
"""

from __future__ import annotations

import hashlib
import json

import audit_reader_docx as audit
import build_reader_docx as builder
from docx import Document
from PIL import Image

ITEMS = [
    {"name": "甲材料", "role": "负载药物"},
    {"name": "乙材料", "role": "靶向肿瘤"},
]


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def make_content(tmp_path, *, reuse_structure_image=False, support=False):
    Image.new("RGB", (80, 80), "red").save(tmp_path / "fig1.png")
    Image.new("RGB", (80, 80), "blue").save(tmp_path / "fig2.png")
    Image.new("RGB", (400, 300), "white").save(tmp_path / "structure.png")
    if support:
        Image.new("RGB", (120, 80), "green").save(tmp_path / "s1.png")
        Image.new("RGB", (120, 80), "yellow").save(tmp_path / "s2.png")
    content = {
        "document_title": "题目与逐图解读",
        "citation": "论文：Author et al., Title, Journal,Year",
        "materials": {
            "items": ITEMS,
            "structure": {
                "label": "Fig.1A",
                "title": "材料结合后的结构示意",
                "image": "fig1.png" if reuse_structure_image else "structure.png",
                "explanation": "各组分通过静电吸附结合为核壳结构。",
            },
            "summary": "本文用核壳结构材料体系以响应性释放的方式解决了递送问题。",
        },
        "field_positioning": {
            "heading": "领域定位与本文解决的问题",
            "paragraphs": ["领域问题"],
        },
        "logic_chain": [
            {"title": "Step1：为什么需要这一步", "text": "对应Fig.1：结论。"}
        ],
        "evidence_blocks": [
            {
                "kind": "figure",
                "label": "Fig.1",
                "title": "Figure 1：主图一",
                "image": "fig1.png",
                "paper_claim": "表述一。",
                "panel_explanations": [
                    {"title": "Fig.1A：读图", "text": "读图。"}
                ],
                "logic_role": "作用一。",
            },
            {
                "kind": "figure",
                "label": "Fig.2",
                "title": "Figure 2：主图二",
                "image": "fig2.png",
                "paper_claim": "表述二。",
                "panel_explanations": [
                    {"title": "Fig.2A：读图", "text": "读图。"}
                ],
                "logic_role": "作用二。",
            },
        ],
    }
    if support:
        content["evidence_blocks"][0]["support_text"] = "补充材料支持主结论。"
        content["evidence_blocks"][0]["support_images"] = [
            {"label": "Fig.S1", "path": "s1.png"},
            {"label": "Fig.S2", "path": "s2.png"},
        ]
    return content


def build(tmp_path, **kwargs):
    content = make_content(tmp_path, **kwargs)
    content_path = tmp_path / "content.json"
    content_path.write_text(json.dumps(content, ensure_ascii=False), encoding="utf-8")
    docx_path = tmp_path / "reader.docx"
    builder.build_document(content, builder.DEFAULT_TEMPLATE, tmp_path, docx_path)
    manifest_path = tmp_path / "image-manifest.json"
    return content, content_path, docx_path, manifest_path


def load_manifest(manifest_path):
    return json.loads(manifest_path.read_text(encoding="utf-8"))


def rewrite_manifest(manifest_path, manifest):
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")


class TestImageManifest:
    def test_build_writes_ordered_image_manifest(self, tmp_path):
        _content, _content_path, docx_path, manifest_path = build(tmp_path)
        manifest = load_manifest(manifest_path)
        assert manifest["schema_version"] == 1
        assert manifest["docx_sha256"] == sha256(docx_path)
        assert [
            (record["kind"], record["object"]) for record in manifest["images"]
        ] == [
            ("materials_structure", "Fig.1A：材料结合后的结构示意"),
            ("evidence_main", "Fig.1"),
            ("evidence_main", "Fig.2"),
        ]
        for record in manifest["images"]:
            assert len(record["embedded_sha256"]) == 64
            assert record["width_emu"] > 0
            assert record["height_emu"] > 0
        for record, filename in zip(manifest["images"][1:], ("fig1.png", "fig2.png")):
            assert record["source"] == str((tmp_path / filename).resolve())
            assert record["source_sha256"] == sha256(tmp_path / filename)

    def test_audit_accepts_a_fresh_image_manifest(self, tmp_path):
        _content, content_path, docx_path, _manifest = build(tmp_path)
        report = audit.audit(docx_path, content_path)
        assert report["errors"] == []

    def test_swapped_same_size_embedded_bytes_are_flagged(self, tmp_path):
        _content, content_path, docx_path, _manifest = build(tmp_path)
        document = Document(docx_path)
        shapes = document.inline_shapes
        relationship_ids = [
            shape._inline.graphic.graphicData.pic.blipFill.blip.embed
            for shape in shapes
        ]
        parts = document.part.related_parts
        first = parts[relationship_ids[1]].blob
        parts[relationship_ids[1]]._blob = parts[relationship_ids[2]].blob
        parts[relationship_ids[2]]._blob = first
        document.save(docx_path)
        report = audit.audit(docx_path, content_path)
        embedded_errors = [
            error for error in report["errors"] if "embedded bytes" in error
        ]
        assert len(embedded_errors) == 2

    def test_same_source_image_may_be_reused(self, tmp_path):
        _content, content_path, docx_path, manifest_path = build(
            tmp_path, reuse_structure_image=True
        )
        manifest = load_manifest(manifest_path)
        sources = [record["source"] for record in manifest["images"]]
        assert sources[0] == sources[1] == str((tmp_path / "fig1.png").resolve())
        report = audit.audit(docx_path, content_path)
        assert report["errors"] == []

    def test_missing_manifest_is_only_a_warning(self, tmp_path):
        _content, content_path, docx_path, manifest_path = build(tmp_path)
        manifest_path.unlink()
        report = audit.audit(docx_path, content_path)
        assert report["errors"] == []
        assert any(
            "image identity unverified" in warning for warning in report["warnings"]
        )

    def test_changed_source_file_is_flagged(self, tmp_path):
        _content, content_path, docx_path, _manifest = build(tmp_path)
        (tmp_path / "fig1.png").write_bytes(b"changed source bytes")
        report = audit.audit(docx_path, content_path)
        assert any("source file changed since build" in error for error in report["errors"])

    def test_removed_embedded_image_is_flagged(self, tmp_path):
        _content, content_path, docx_path, _manifest = build(tmp_path)
        document = Document(docx_path)
        picture_paragraphs = [
            paragraph
            for paragraph in document.paragraphs
            if paragraph.style.name == "FER Picture"
        ]
        picture_paragraphs[-1]._element.getparent().remove(
            picture_paragraphs[-1]._element
        )
        document.save(docx_path)
        report = audit.audit(docx_path, content_path)
        assert any(
            "Image manifest records 3 != embedded image occurrences 2" in error
            for error in report["errors"]
        )

    def test_manifest_order_mismatch_is_flagged(self, tmp_path):
        _content, content_path, docx_path, manifest_path = build(tmp_path)
        manifest = load_manifest(manifest_path)
        manifest["images"][0], manifest["images"][1] = (
            manifest["images"][1],
            manifest["images"][0],
        )
        rewrite_manifest(manifest_path, manifest)
        report = audit.audit(docx_path, content_path)
        assert any(
            "Image manifest order does not match" in error for error in report["errors"]
        )

    def test_unsupported_manifest_schema_is_flagged(self, tmp_path):
        _content, content_path, docx_path, manifest_path = build(tmp_path)
        manifest = load_manifest(manifest_path)
        manifest["schema_version"] = 99
        rewrite_manifest(manifest_path, manifest)
        report = audit.audit(docx_path, content_path)
        assert any("unsupported schema" in error for error in report["errors"])

    def test_corrupt_manifest_is_flagged(self, tmp_path):
        _content, content_path, docx_path, manifest_path = build(tmp_path)
        manifest_path.write_text("{broken", encoding="utf-8")
        report = audit.audit(docx_path, content_path)
        assert any("cannot be read" in error for error in report["errors"])


class TestSupportSheetIdentity:
    def test_support_sheet_members_are_recorded_and_audited(self, tmp_path):
        _content, content_path, docx_path, manifest_path = build(
            tmp_path, support=True
        )
        manifest = load_manifest(manifest_path)
        assert [record["kind"] for record in manifest["images"]] == [
            "materials_structure",
            "evidence_main",
            "support_sheet",
            "evidence_main",
        ]
        support_record = manifest["images"][2]
        assert support_record["kind"] == "support_sheet"
        assert support_record["source"] is None
        assert [member["path"] for member in support_record["members"]] == [
            str((tmp_path / "s1.png").resolve()),
            str((tmp_path / "s2.png").resolve()),
        ]
        assert support_record["members"][0]["sha256"] == sha256(tmp_path / "s1.png")
        report = audit.audit(docx_path, content_path)
        assert report["errors"] == []

    def test_missing_support_member_is_flagged(self, tmp_path):
        _content, content_path, docx_path, _manifest = build(tmp_path, support=True)
        (tmp_path / "s1.png").unlink()
        report = audit.audit(docx_path, content_path)
        assert any("support member 1 is missing" in error for error in report["errors"])

    def test_changed_support_member_is_flagged(self, tmp_path):
        _content, content_path, docx_path, _manifest = build(tmp_path, support=True)
        Image.new("RGB", (120, 80), "black").save(tmp_path / "s2.png")
        report = audit.audit(docx_path, content_path)
        assert any(
            "support member 2 changed since build" in error
            for error in report["errors"]
        )
