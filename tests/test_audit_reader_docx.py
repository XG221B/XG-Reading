import json
from copy import deepcopy

import audit_reader_docx as audit
import build_reader_docx as builder
import pytest
from docx import Document
from docx.shared import Cm
from PIL import Image


def make_content(tmp_path):
    image = tmp_path / "fig1.png"
    Image.new("RGB", (80, 80), "white").save(image)
    return {
        "document_title": "题目与逐图解读",
        "citation": "论文：Author et al., Title, Journal,Year",
        "field_positioning": {
            "heading": "领域定位与本文解决的问题",
            "paragraphs": ["领域问题", "具体缺口", "本文贡献与边界"],
        },
        "logic_chain": [
            {
                "title": "Step1：为什么需要这一步",
                "text": "对应Fig.1：实验、结果及其支持的结论。",
            }
        ],
        "evidence_blocks": [
            {
                "kind": "figure",
                "label": "Fig.1",
                "title": "Figure 1：主图英文标题的准确中文翻译",
                "image": "fig1.png",
                "paper_claim": "作者对该证据块的主要表述。",
                "panel_explanations": [
                    {
                        "title": "Fig.1A：该panel图注内容的中文译名",
                        "text": "Fig.1A的读图方法与意义。",
                    },
                    {
                        "title": "Fig.1B：第二个panel的中文译名",
                        "text": "Fig.1B的读图方法与意义。",
                    },
                ],
                "logic_role": "该证据块在全文论证中的作用。",
            }
        ],
    }


@pytest.fixture
def built(tmp_path):
    content = make_content(tmp_path)
    content_path = tmp_path / "content.json"
    content_path.write_text(
        json.dumps(content, ensure_ascii=False), encoding="utf-8"
    )
    docx_path = tmp_path / "reader.docx"
    builder.build_document(
        content,
        builder.DEFAULT_TEMPLATE,
        content_path.parent,
        docx_path,
    )
    return content, content_path, docx_path


class TestRoundTrip:
    def test_built_docx_passes_audit(self, built):
        _, content_path, docx_path = built
        report = audit.audit(docx_path, content_path)
        assert report["errors"] == []

    def test_audit_without_content_does_not_crash(self, built):
        _, _, docx_path = built
        report = audit.audit(docx_path, None)
        assert isinstance(report["errors"], list)

    def test_blank_docx_does_not_crash_audit(self, tmp_path):
        blank = tmp_path / "blank.docx"
        Document().save(blank)
        report = audit.audit(blank, None)
        assert isinstance(report["errors"], list)
        assert any("Missing style" in error for error in report["errors"])


class TestStyleContract:
    def test_line_spacing_deviation_flagged(self, built):
        _, _, docx_path = built
        document = Document(docx_path)
        document.styles["FER Body"].paragraph_format.line_spacing = 1.5
        document.save(docx_path)
        report = audit.audit(docx_path, None)
        assert any("line spacing" in error for error in report["errors"])

    def test_alignment_deviation_flagged(self, built):
        _, _, docx_path = built
        document = Document(docx_path)
        document.styles["FER Step"].paragraph_format.alignment = 2
        document.save(docx_path)
        report = audit.audit(docx_path, None)
        assert any("alignment" in error for error in report["errors"])

    def test_bold_deviation_flagged(self, built):
        _, _, docx_path = built
        document = Document(docx_path)
        document.styles["FER Title"].font.bold = False
        document.save(docx_path)
        report = audit.audit(docx_path, None)
        assert any("bold" in error for error in report["errors"])


class TestParagraphSequence:
    def test_boundary_space_flagged(self, built):
        _, _, docx_path = built
        document = Document(docx_path)
        for paragraph in document.paragraphs:
            if "读图方法" in paragraph.text:
                paragraph.runs[0].text = "细胞 A 的荧光。"
                break
        document.save(docx_path)
        report = audit.audit(docx_path, None)
        assert any("boundary space" in error for error in report["errors"])

    def test_dropped_paragraph_flagged(self, built):
        _, content_path, docx_path = built
        document = Document(docx_path)
        for paragraph in list(document.paragraphs):
            if "Fig.1B的读图方法" in paragraph.text:
                paragraph._element.getparent().remove(paragraph._element)
                break
        document.save(docx_path)
        report = audit.audit(docx_path, content_path)
        assert any("text" in error for error in report["errors"])

    def test_spacer_after_paper_claim_flagged(self, built):
        _, content_path, docx_path = built
        document = Document(docx_path)
        paragraphs = document.paragraphs
        spacer = next(
            paragraph
            for paragraph in paragraphs
            if not paragraph.text.strip()
            and paragraph.paragraph_format.line_spacing == 0.2
        )
        paper_claim = next(
            paragraph
            for paragraph in paragraphs
            if "作者对该证据块" in paragraph.text
        )
        paper_claim._element.addnext(deepcopy(spacer._element))
        document.save(docx_path)
        report = audit.audit(docx_path, content_path)
        assert any("spacer" in error for error in report["errors"])


class TestImageWidths:
    def test_wrong_main_image_width_flagged(self, built):
        _, content_path, docx_path = built
        document = Document(docx_path)
        document.inline_shapes[0].width = int(Cm(10))
        document.save(docx_path)
        report = audit.audit(docx_path, content_path)
        assert any("width" in error for error in report["errors"])


class TestNewContractChecks:
    def test_page_break_before_on_non_figure_style_flagged(self, built):
        _, _, docx_path = built
        document = Document(docx_path)
        document.styles["FER Body"].paragraph_format.page_break_before = True
        document.save(docx_path)
        report = audit.audit(docx_path, None)
        assert any("page break" in error.lower() for error in report["errors"])

    def test_header_text_flagged(self, built):
        _, _, docx_path = built
        document = Document(docx_path)
        document.sections[0].header.paragraphs[0].add_run("页眉文字")
        document.save(docx_path)
        report = audit.audit(docx_path, None)
        assert any("header" in error.lower() for error in report["errors"])

    def test_unbolded_label_run_flagged(self, built):
        _, _, docx_path = built
        document = Document(docx_path)
        for paragraph in document.paragraphs:
            if "Fig.1总结" in paragraph.text:
                paragraph.runs[0].bold = False
                break
        document.save(docx_path)
        report = audit.audit(docx_path, None)
        assert any("run" in error.lower() for error in report["errors"])

    def test_bolded_body_run_flagged(self, built):
        _, _, docx_path = built
        document = Document(docx_path)
        for paragraph in document.paragraphs:
            if "Fig.1总结" in paragraph.text:
                paragraph.runs[1].bold = True
                break
        document.save(docx_path)
        report = audit.audit(docx_path, None)
        assert any("run" in error.lower() for error in report["errors"])
